from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "17365D"
WHITE = "FFFFFF"
PALE_BLUE = "EAF1F7"
PALE_GREEN = "E2F0D9"
TEXT_COLOR = "243247"
MUTED_COLOR = "66788A"

PERCENT_FORMAT = "0.00%"
P_VALUE_FORMAT = "0.000000"


def style_header(worksheet, row_number=1):
    for cell in worksheet[row_number]:
        if cell.value is None:
            continue

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=NAVY,
        )
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=WHITE,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[row_number].height = 32


def style_body(worksheet, first_row=2):
    for row in worksheet.iter_rows(min_row=first_row):
        for cell in row:
            cell.font = Font(
                name="Calibri",
                size=11,
                color=TEXT_COLOR,
            )
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if cell.row % 2 == 0:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=PALE_BLUE,
                )

        if worksheet.row_dimensions[row[0].row].height is None:
            worksheet.row_dimensions[row[0].row].height = 25


def set_column_widths(worksheet, widths):
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def format_percentage(value):
    if value is None:
        return "N/A"

    return f"{value:.2%}"


def get_variant_details(result):
    variants = result.get("variants")

    if not isinstance(variants, list) or not variants:
        raise ValueError(
            "Experiment result must contain variant details."
        )

    return variants


def get_control_comparisons(result):
    return {
        comparison["variant_b"]: comparison
        for comparison in result["comparisons"]
        if comparison["variant_a"] == "A"
    }


def get_best_improvement(result):
    candidates = []

    for comparison in result["comparisons"]:
        difference = (
                comparison["success_rate_b"]
                - comparison["success_rate_a"]
        )

        if comparison["is_significant"] and difference > 0:
            candidates.append(comparison)

    if not candidates:
        return None

    return max(
        candidates,
        key=lambda comparison: (
                comparison["success_rate_b"]
                - comparison["success_rate_a"]
        ),
    )


def get_comparison_label(comparison):
    return (
        f'{comparison["variant_a"]} vs '
        f'{comparison["variant_b"]}'
    )


def get_improvement_label(comparison):
    return (
        f'{comparison["variant_b"]} vs '
        f'{comparison["variant_a"]}'
    )


def get_statistical_conclusion(comparison):
    if not comparison["is_significant"]:
        return "Not significant"

    difference = (
            comparison["success_rate_b"]
            - comparison["success_rate_a"]
    )

    if difference > 0:
        return "Significant increase"

    if difference < 0:
        return "Significant decrease"

    return "No difference"


def build_interpretation(result, best):
    comparisons = result["comparisons"]

    if best is None:
        significant_count = sum(
            comparison["is_significant"]
            for comparison in comparisons
        )

        if significant_count:
            return (
                "No statistically significant improvement "
                "was identified. Some comparisons showed "
                "a statistically significant decrease."
            )

        return (
            "No statistically significant difference "
            "was identified among the tested comparisons. "
            "Observed success rate differences are not "
            "sufficient to establish an improvement."
        )

    variant_a = best["variant_a"]
    variant_b = best["variant_b"]

    rate_a = best["success_rate_a"]
    rate_b = best["success_rate_b"]

    difference = (rate_b - rate_a) * 100

    if rate_a == 0:
        relative_text = (
            "Relative improvement is undefined "
            "because the baseline success rate is zero"
        )
    else:
        relative_improvement = (
                (rate_b - rate_a) / rate_a
        )

        relative_text = (
            f"equivalent to a "
            f"{format_percentage(relative_improvement)} "
            f"relative improvement"
        )

    interpretation = (
        f"Variant {variant_b} achieved a "
        f"{format_percentage(rate_b)} success rate "
        f"compared with {format_percentage(rate_a)} "
        f"for Variant {variant_a}. "
        f"The difference is {difference:.2f} "
        f"percentage points, {relative_text}, "
        f"and is statistically significant."
    )

    other_comparisons = [
        comparison
        for comparison in comparisons
        if comparison is not best
    ]

    not_significant = [
        comparison
        for comparison in other_comparisons
        if not comparison["is_significant"]
    ]

    if not_significant:
        labels = ", ".join(
            get_comparison_label(comparison)
            for comparison in not_significant
        )

        interpretation += (
            f" The following comparisons were "
            f"not statistically significant: {labels}."
        )

    return interpretation


def create_summary_sheet(workbook, result):
    worksheet = workbook.active
    worksheet.title = "Summary"

    variants = get_variant_details(result)
    best = get_best_improvement(result)

    control = next(
        variant
        for variant in variants
        if variant["name"] == "A"
    )

    comparison_labels = ", ".join(
        get_comparison_label(comparison)
        for comparison in result["comparisons"]
    )

    comparison_mode = (
        "Control vs alternatives"
        if result["comparison_mode"] == "control"
        else "All pairwise comparisons"
    )

    significant_count = sum(
        comparison["is_significant"]
        for comparison in result["comparisons"]
    )

    rows = [
        ("Result overview", "Value"),
        (
            "Best statistically supported improvement",
            (
                get_improvement_label(best)
                if best is not None
                else "None identified"
            ),
        ),
        ("Control variant", "A"),
        ("Comparisons", comparison_labels),
        ("Comparison mode", comparison_mode),
        ("Significant comparisons", significant_count),
        (
            "Control success rate",
            control["success_rate"],
        ),
    ]

    if best is not None:
        rows.extend([
            (
                f'Variant {best["variant_b"]} success rate',
                best["success_rate_b"],
            ),
            (
                "Success rate difference",
                (
                        best["success_rate_b"]
                        - best["success_rate_a"]
                ),
            ),
            (
                "Relative improvement",
                (
                    (
                            best["success_rate_b"]
                            - best["success_rate_a"]
                    ) / best["success_rate_a"]
                    if best["success_rate_a"] != 0
                    else None
                ),
            ),
        ])
    else:
        rows.extend([
            ("Improved variant success rate", None),
            ("Success rate difference", None),
            ("Relative improvement", None),
        ])

    rows.append((
        "Interpretation",
        build_interpretation(result, best),
    ))

    for row in rows:
        worksheet.append(row)

    style_body(worksheet)
    style_header(worksheet)

    set_column_widths(
        worksheet,
        {
            "A": 43,
            "B": 65,
        },
    )

    for row_number in (7, 8, 9, 10):
        worksheet.cell(
            row=row_number,
            column=2,
        ).number_format = PERCENT_FORMAT

    for cell in worksheet[2]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=PALE_GREEN,
        )
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=TEXT_COLOR,
        )

    interpretation_row = len(rows)

    label_cell = worksheet.cell(
        row=interpretation_row,
        column=1,
    )
    text_cell = worksheet.cell(
        row=interpretation_row,
        column=2,
    )

    label_cell.font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color=TEXT_COLOR,
    )

    text_cell.font = Font(
        name="Calibri",
        size=11,
        bold=False,
        color=TEXT_COLOR,
    )

    text_cell.alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    label_cell.alignment = Alignment(
        vertical="top",
    )

    worksheet.row_dimensions[
        interpretation_row
    ].height = max(
        75,
        16 * (
                len(str(text_cell.value)) // 65 + 2
        ),
    )

    worksheet.freeze_panes = None
    worksheet.sheet_view.showGridLines = False


def create_comparisons_sheet(workbook, result):
    worksheet = workbook.create_sheet(
        "Comparison Details"
    )

    headers = [
        "Variant",
        "Role",
        "Sample size",
        "Success count",
        "Success rate",
        "Improvement vs A",
        "Adjusted p-value",
        "Statistical conclusion",
    ]

    worksheet.append(headers)

    variants = get_variant_details(result)
    control_comparisons = get_control_comparisons(result)

    control = next(
        variant
        for variant in variants
        if variant["name"] == "A"
    )

    control_rate = control["success_rate"]

    for variant in variants:
        name = variant["name"]

        if name == "A":
            worksheet.append([
                name,
                "Control",
                variant["sample_size"],
                variant["success_count"],
                variant["success_rate"],
                None,
                None,
                "Reference",
            ])
            continue

        comparison = control_comparisons.get(name)

        if control_rate == 0:
            improvement = None
        else:
            improvement = (
                                  variant["success_rate"] - control_rate
                          ) / control_rate

        worksheet.append([
            name,
            "Alternative",
            variant["sample_size"],
            variant["success_count"],
            variant["success_rate"],
            improvement,
            (
                comparison["adjusted_p_value"]
                if comparison is not None
                else None
            ),
            (
                get_statistical_conclusion(comparison)
                if comparison is not None
                else "Not available"
            ),
        ])

    style_body(worksheet)
    style_header(worksheet)

    for row_number in range(2, len(variants) + 2):
        for column in (5, 6):
            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = PERCENT_FORMAT

        worksheet.cell(
            row=row_number,
            column=7,
        ).number_format = P_VALUE_FORMAT

    set_column_widths(
        worksheet,
        {
            "A": 15,
            "B": 18,
            "C": 18,
            "D": 19,
            "E": 19,
            "F": 23,
            "G": 22,
            "H": 27,
        },
    )

    # In all-pairs mode, comparisons between alternatives
    # are shown separately to preserve the one-row-per-variant
    # structure of the main table.
    if result["comparison_mode"] == "all_pairs":
        extra_comparisons = [
            comparison
            for comparison in result["comparisons"]
            if comparison["variant_a"] != "A"
        ]

        if extra_comparisons:
            start_row = len(variants) + 4

            worksheet.cell(
                row=start_row,
                column=1,
                value="Additional pairwise comparisons",
            )

            worksheet.cell(
                row=start_row,
                column=1,
            ).font = Font(
                name="Calibri",
                bold=True,
                size=12,
                color=NAVY,
            )

            header_row = start_row + 1

            extra_headers = [
                "Comparison",
                "Success rate A",
                "Success rate B",
                "Difference",
                "Adjusted p-value",
                "Statistical conclusion",
            ]

            for column, header in enumerate(
                    extra_headers,
                    start=1,
            ):
                worksheet.cell(
                    row=header_row,
                    column=column,
                    value=header,
                )

            for comparison in extra_comparisons:
                worksheet.append([
                    get_comparison_label(comparison),
                    comparison["success_rate_a"],
                    comparison["success_rate_b"],
                    (
                            comparison["success_rate_b"]
                            - comparison["success_rate_a"]
                    ),
                    comparison["adjusted_p_value"],
                    get_statistical_conclusion(comparison),
                ])

            for row in worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=worksheet.max_row,
                    max_col=6,
            ):
                for cell in row:
                    cell.font = Font(
                        name="Calibri",
                        size=11,
                        color=TEXT_COLOR,
                    )
                    cell.alignment = Alignment(
                        vertical="center",
                    )

                for index in (1, 2, 3):
                    row[index].number_format = (
                        PERCENT_FORMAT
                    )

                row[4].number_format = P_VALUE_FORMAT

            style_header(worksheet, header_row)

    worksheet.freeze_panes = None
    worksheet.auto_filter.ref = (
        f"A1:H{len(variants) + 1}"
    )
    worksheet.sheet_view.showGridLines = False


def create_guide_sheet(workbook):
    worksheet = workbook.create_sheet("Guide")

    rows = [
        ("Metric", "Meaning"),
        (
            "Success rate",
            "Success count divided by sample size.",
        ),
        (
            "Improvement vs A",
            "Relative change in success rate "
            "compared with control A.",
        ),
        (
            "Success rate difference",
            "Difference between two success rates; "
            "shown as a percentage, equivalent to "
            "percentage-point change.",
        ),
        (
            "Relative improvement",
            "How much higher the success rate is "
            "relative to the baseline rate.",
        ),
        (
            "Adjusted p-value",
            "Statistical evidence after accounting "
            "for multiple comparisons. Smaller values "
            "indicate stronger evidence against "
            "equal success rates.",
        ),
        (
            "Statistical significance",
            "Whether the observed difference passes "
            "the configured statistical threshold.",
        ),
        (
            "Reference",
            "The control variant used as the baseline.",
        ),
    ]

    for row in rows:
        worksheet.append(row)

    style_body(worksheet)
    style_header(worksheet)

    set_column_widths(
        worksheet,
        {
            "A": 30,
            "B": 85,
        },
    )

    for row_number in range(2, len(rows) + 1):
        worksheet.row_dimensions[row_number].height = 34

    worksheet.freeze_panes = None
    worksheet.sheet_view.showGridLines = False


def generate_excel_report(
        experiment_result: dict,
        output_path: str | Path,
) -> Path:
    path = Path(output_path)

    if path.suffix.lower() != ".xlsx":
        raise ValueError(
            "Excel report path must end with .xlsx."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    try:
        create_summary_sheet(
            workbook,
            experiment_result,
        )
        create_comparisons_sheet(
            workbook,
            experiment_result,
        )
        create_guide_sheet(workbook)

        workbook.save(path)
    finally:
        workbook.close()

    return path
