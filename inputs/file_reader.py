
import csv
from pathlib import Path

from openpyxl import load_workbook

from core.models import BinaryVariant
from core.validator import validate_experiment


REQUIRED_COLUMNS = (
    "variant",
    "sample_size",
    "success_count",
)


def parse_integer(value, column_name: str) -> int:
    if isinstance(value, bool):
        raise ValueError(
            f"Column '{column_name}' must contain an integer."
        )

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(
                f"Column '{column_name}' must contain an integer."
            )
        return int(value)

    if isinstance(value, str):
        value = value.strip()

        if value.isdecimal():
            return int(value)

    raise ValueError(
        f"Column '{column_name}' must contain an integer."
    )


def parse_rows(rows: list[dict]) -> list[BinaryVariant]:
    if not rows:
        raise ValueError("Input file contains no data.")

    variants = []

    for row_number, row in enumerate(rows, start=2):
        missing_columns = [
            column
            for column in REQUIRED_COLUMNS
            if column not in row
        ]

        if missing_columns:
            raise ValueError(
                f"Missing required columns: {missing_columns}"
            )

        try:
            name = row["variant"]

            if not isinstance(name, str):
                raise ValueError(
                    "Column 'variant' must contain text."
                )

            variant = BinaryVariant(
                name=name.strip(),
                sample_size=parse_integer(
                    row["sample_size"],
                    "sample_size",
                ),
                success_count=parse_integer(
                    row["success_count"],
                    "success_count",
                ),
            )

            variants.append(variant)

        except ValueError as error:
            raise ValueError(
                f"Invalid data in row {row_number}: {error}"
            ) from error

    validate_experiment(variants)

    return variants


def read_csv(file_path: Path) -> list[BinaryVariant]:
    with file_path.open(
        mode="r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        if reader.fieldnames is None:
            raise ValueError("CSV file is empty.")

        missing_columns = [
            column
            for column in REQUIRED_COLUMNS
            if column not in reader.fieldnames
        ]

        if missing_columns:
            raise ValueError(
                f"Missing required columns: {missing_columns}"
            )

        rows = list(reader)

    return parse_rows(rows)


def read_excel(file_path: Path) -> list[BinaryVariant]:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)

        headers = next(values, None)

        if headers is None:
            raise ValueError("Excel file is empty.")

        missing_columns = [
            column
            for column in REQUIRED_COLUMNS
            if column not in headers
        ]

        if missing_columns:
            raise ValueError(
                f"Missing required columns: {missing_columns}"
            )

        rows = [
            dict(zip(headers, values_row))
            for values_row in values
            if any(value is not None for value in values_row)
        ]

        return parse_rows(rows)

    finally:
        workbook.close()


def read_experiment_file(
    file_path: str,
) -> list[BinaryVariant]:

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".csv":
        return read_csv(path)

    if extension == ".xlsx":
        return read_excel(path)

    raise ValueError(
        "Unsupported file format. Use CSV or XLSX."
    )
