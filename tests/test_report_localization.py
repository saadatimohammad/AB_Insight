
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from core.engine import run_experiment
from core.interpreter import interpret_experiment
from core.models import BinaryVariant
from reports.console import generate_console_report
from reports.excel import generate_excel_report
from reports.locales import en


class TestReportLocalization(unittest.TestCase):

    def setUp(self):
        variants = [
            BinaryVariant("A", 1000, 100),
            BinaryVariant("B", 1000, 200),
        ]

        result = run_experiment(
            variants=variants,
            comparison_mode="control",
            alpha=0.05,
        )

        self.result = interpret_experiment(result)

    def test_translation_sections_exist(self):
        sections = [
            "CONSOLE",
            "INTERPRETATIONS",
            "EXCEL",
            "GUIDE",
            "SUMMARY_INTERPRETATIONS",
        ]

        for section in sections:
            with self.subTest(section=section):
                translations = getattr(en, section)

                self.assertIsInstance(
                    translations,
                    dict,
                )

                self.assertTrue(translations)

    def test_translation_values_are_strings(self):
        sections = [
            en.CONSOLE,
            en.INTERPRETATIONS,
            en.EXCEL,
            en.GUIDE,
            en.SUMMARY_INTERPRETATIONS,
        ]

        for translations in sections:
            for key, value in translations.items():
                with self.subTest(key=key):
                    self.assertIsInstance(key, str)
                    self.assertIsInstance(value, str)
                    self.assertTrue(value.strip())

    def test_summary_translation_placeholders(self):
        required_placeholders = {
            "best_improvement": {
                "variant_a",
                "variant_b",
                "rate_a",
                "rate_b",
                "difference",
                "relative_text",
            },
            "relative_improvement": {
                "value",
            },
            "other_not_significant": {
                "comparisons",
            },
        }

        from string import Formatter

        formatter = Formatter()

        for key, expected in required_placeholders.items():
            template = en.SUMMARY_INTERPRETATIONS[key]

            actual = {
                field_name
                for _, field_name, _, _ in formatter.parse(
                    template
                )
                if field_name is not None
            }

            with self.subTest(key=key):
                self.assertEqual(actual, expected)

    def test_console_uses_translated_labels(self):
        with patch.dict(
            en.CONSOLE,
            {"comparison_mode": "TEST MODE LABEL"},
        ):
            report = generate_console_report(
                self.result,
                language="en",
            )

        self.assertIn(
            "TEST MODE LABEL:",
            report,
        )

    def test_console_uses_translated_interpretation(self):
        with patch.dict(
            en.INTERPRETATIONS,
            {
                "significant_increase": (
                    "TEST SIGNIFICANT INCREASE"
                ),
            },
        ):
            report = generate_console_report(
                self.result,
                language="en",
            )

        self.assertIn(
            "TEST SIGNIFICANT INCREASE",
            report,
        )

    def test_excel_uses_translated_content(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = (
                Path(directory) / "localized_report.xlsx"
            )

            with patch.dict(
                en.EXCEL,
                {
                    "summary_sheet": "Test Summary",
                    "result_overview": "TEST OVERVIEW",
                },
            ):
                with patch.dict(
                    en.GUIDE,
                    {
                        "success_rate": (
                            "TEST SUCCESS RATE DESCRIPTION"
                        ),
                    },
                ):
                    generate_excel_report(
                        self.result,
                        output_path,
                        language="en",
                    )

            workbook = load_workbook(output_path)

            try:
                self.assertIn(
                    "Test Summary",
                    workbook.sheetnames,
                )

                summary = workbook["Test Summary"]

                self.assertEqual(
                    summary["A1"].value,
                    "TEST OVERVIEW",
                )

                guide = workbook["Guide"]

                self.assertEqual(
                    guide["B2"].value,
                    "TEST SUCCESS RATE DESCRIPTION",
                )
            finally:
                workbook.close()

    def test_excel_uses_translated_interpretation(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = (
                Path(directory) / "interpretation.xlsx"
            )

            with patch.dict(
                en.SUMMARY_INTERPRETATIONS,
                {
                    "best_improvement": (
                        "TEST BEST IMPROVEMENT "
                        "{variant_a} {variant_b} "
                        "{rate_a} {rate_b} "
                        "{difference} {relative_text}"
                    ),
                },
            ):
                generate_excel_report(
                    self.result,
                    output_path,
                    language="en",
                )

            workbook = load_workbook(output_path)

            try:
                summary = workbook["Summary"]

                interpretation = summary["B11"].value

                self.assertIn(
                    "TEST BEST IMPROVEMENT",
                    interpretation,
                )
            finally:
                workbook.close()

    def test_reports_reject_unsupported_language(self):
        with self.assertRaises(ValueError):
            generate_console_report(
                self.result,
                language="de",
            )

        with tempfile.TemporaryDirectory() as directory:
            output_path = (
                Path(directory) / "invalid_language.xlsx"
            )

            with self.assertRaises(ValueError):
                generate_excel_report(
                    self.result,
                    output_path,
                    language="de",
                )

            self.assertFalse(output_path.exists())


if __name__ == "__main__":
    unittest.main()
