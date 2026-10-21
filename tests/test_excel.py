
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.engine import run_experiment
from core.interpreter import interpret_experiment
from core.models import BinaryVariant
from reports.excel import generate_excel_report


class TestExcelReport(unittest.TestCase):

    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_directory.cleanup)

        self.output_path = (
            Path(self.temp_directory.name) / "report.xlsx"
        )

        self.variants = [
            BinaryVariant("A", 1000, 120),
            BinaryVariant("B", 1000, 150),
            BinaryVariant("C", 1000, 180),
        ]

    def create_report(
        self,
        variants=None,
        comparison_mode="control",
    ):
        if variants is None:
            variants = self.variants

        result = run_experiment(
            variants=variants,
            comparison_mode=comparison_mode,
        )

        interpreted_result = interpret_experiment(result)

        generate_excel_report(
            interpreted_result,
            self.output_path,
        )

        workbook = load_workbook(self.output_path)
        self.addCleanup(workbook.close)

        return workbook

    def test_excel_file_created(self):
        self.create_report()

        self.assertTrue(self.output_path.is_file())

    def test_expected_sheets(self):
        workbook = self.create_report()

        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Comparison Details", "Guide"],
        )

    def test_variant_input_data(self):
        workbook = self.create_report()
        worksheet = workbook["Comparison Details"]

        self.assertEqual(worksheet["A2"].value, "A")
        self.assertEqual(worksheet["C2"].value, 1000)
        self.assertEqual(worksheet["D2"].value, 120)

        self.assertEqual(worksheet["A4"].value, "C")
        self.assertEqual(worksheet["C4"].value, 1000)
        self.assertEqual(worksheet["D4"].value, 180)

    def test_success_rates_are_numeric(self):
        workbook = self.create_report()
        worksheet = workbook["Comparison Details"]

        self.assertAlmostEqual(
            worksheet["E2"].value,
            0.12,
        )
        self.assertAlmostEqual(
            worksheet["E4"].value,
            0.18,
        )

        self.assertEqual(
            worksheet["E2"].number_format,
            "0.00%",
        )

        self.assertIsInstance(
            worksheet["E2"].value,
            (int, float),
        )

    def test_best_significant_improvement(self):
        workbook = self.create_report()
        worksheet = workbook["Summary"]

        self.assertEqual(
            worksheet["B2"].value,
            "C vs A",
        )

        self.assertAlmostEqual(
            worksheet["B7"].value,
            0.12,
        )
        self.assertAlmostEqual(
            worksheet["B8"].value,
            0.18,
        )
        self.assertAlmostEqual(
            worksheet["B9"].value,
            0.06,
        )
        self.assertAlmostEqual(
            worksheet["B10"].value,
            0.50,
        )

    def test_no_significant_improvement(self):
        variants = [
            BinaryVariant("A", 1000, 120),
            BinaryVariant("B", 1000, 121),
            BinaryVariant("C", 1000, 122),
        ]

        workbook = self.create_report(variants)
        worksheet = workbook["Summary"]

        self.assertEqual(
            worksheet["B2"].value,
            "None identified",
        )

        self.assertIsNone(worksheet["B8"].value)
        self.assertIsNone(worksheet["B9"].value)
        self.assertIsNone(worksheet["B10"].value)

    def test_all_pairs_mode(self):
        workbook = self.create_report(
            comparison_mode="all_pairs",
        )

        worksheet = workbook["Comparison Details"]

        self.assertEqual(
            worksheet["A8"].value,
            "Comparison",
        )

        self.assertEqual(
            worksheet["A9"].value,
            "B vs C",
        )

    def test_no_frozen_panes(self):
        workbook = self.create_report()

        for worksheet in workbook.worksheets:
            with self.subTest(sheet=worksheet.title):
                self.assertIsNone(
                    worksheet.freeze_panes
                )

    def test_invalid_output_extension(self):
        result = interpret_experiment(
            run_experiment(self.variants)
        )

        invalid_path = (
            Path(self.temp_directory.name) / "report.csv"
        )

        with self.assertRaises(ValueError):
            generate_excel_report(
                result,
                invalid_path,
            )


if __name__ == "__main__":
    unittest.main()
